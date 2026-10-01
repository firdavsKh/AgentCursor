#!/usr/bin/env python3
"""Prepare star-schema model pieces from PB archive files.

This script extracts PB workbooks and report configs from RAR archives,
normalizes heterogeneous tables into long-form fact rows, and emits dimension
and reference CSVs that match the drill model discussed in documentation.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import shutil
import subprocess
import sys
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


PERIOD_BUCKETS = {
    "current_year",
    "prevent_year",
    "current_cis",
    "current_far",
    "prevent_cis",
    "prevent_far",
}


def enforce_current_year_for_sparse_periods(
    fact_rows: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """Force current_year when period exists without a current/prevent pair.

    Rule:
    - detect tables that have at least one non-empty period_bucket
    - if a table does not contain both current_year and prevent_year buckets,
      then all its non-empty period buckets are normalized to current_year
      and corresponding measure_type_key is also normalized to current_year
    """

    buckets_by_table: dict[int, set[str]] = defaultdict(set)
    for row in fact_rows:
        period_bucket = str(row.get("period_bucket") or "").strip()
        if not period_bucket:
            continue
        pb_table = int(row.get("pb_table") or 0)
        buckets_by_table[pb_table].add(period_bucket)

    enforced_rows: list[dict[str, Any]] = []
    enforce_tables: set[int] = set()
    required_pair = {"current_year", "prevent_year"}

    for pb_table in sorted(buckets_by_table):
        buckets = buckets_by_table[pb_table]
        if buckets and not required_pair.issubset(buckets):
            enforce_tables.add(pb_table)
            enforced_rows.append(
                {
                    "pb_table": pb_table,
                    "original_period_buckets": ",".join(sorted(buckets)),
                    "enforced_period_bucket": "current_year",
                }
            )

    if not enforce_tables:
        return enforced_rows

    for row in fact_rows:
        pb_table = int(row.get("pb_table") or 0)
        period_bucket = str(row.get("period_bucket") or "").strip()
        if pb_table in enforce_tables and period_bucket:
            row["period_bucket"] = "current_year"
            row["measure_type_key"] = "current_year"

    return enforced_rows


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Extract PB archives and prepare fact/dimension model pieces."
    )
    parser.add_argument(
        "--output-rar",
        default="output_files.rar",
        help="Path to output_files.rar archive",
    )
    parser.add_argument(
        "--config-rar",
        default="report_configs.rar",
        help="Path to report_configs.rar archive",
    )
    parser.add_argument(
        "--output-dir",
        default="prepared_model",
        help="Directory where CSV outputs will be written",
    )
    parser.add_argument(
        "--work-dir",
        default=".model_work",
        help="Scratch directory for extraction and temporary artifacts",
    )
    parser.add_argument(
        "--skip-extract",
        action="store_true",
        help="Skip extraction and use already extracted folders in work-dir",
    )
    parser.add_argument(
        "--keep-work-dir",
        action="store_true",
        help="Do not delete/recreate work-dir before processing",
    )
    return parser.parse_args()


def require_utility(binary: str) -> None:
    if shutil.which(binary):
        return
    raise RuntimeError(
        f"Required utility '{binary}' is not available. "
        "Install it (e.g. apt install unar) or provide --skip-extract with "
        "pre-extracted files."
    )


def run_unar(archive_path: Path, target_dir: Path) -> tuple[int, str]:
    cmd = [
        "unar",
        "-quiet",
        "-output-directory",
        str(target_dir),
        str(archive_path),
    ]
    proc = subprocess.run(cmd, capture_output=True, text=True, check=False)
    combined_output = (proc.stdout or "") + (proc.stderr or "")
    return proc.returncode, combined_output


def parse_period_from_filename(name: str) -> tuple[str, int, int | None, str]:
    years = re.findall(r"(20\d{2})", name)
    year = int(max(years)) if years else -1
    quarter: int | None = None
    quarter_label = ""

    if "(I+II+III)" in name:
        quarter = 3
        quarter_label = "I+II+III"
    elif "(I+II)" in name or "(I+I I)" in name:
        quarter = 2
        quarter_label = "I+II"
    else:
        match = re.search(r"_(\d)(?:\D|$)", name)
        if match:
            quarter = int(match.group(1))
            quarter_label = f"Q{quarter}"
        else:
            any_digit = re.search(r"\b(\d)\b", name)
            if any_digit:
                quarter = int(any_digit.group(1))
                quarter_label = f"Q{quarter}"

    if year < 0:
        label = "unknown"
        sort_key = -1
    else:
        label = f"{year}" + (f"-{quarter_label}" if quarter_label else "")
        sort_key = year * 10 + (quarter if quarter is not None else 0)

    return label, sort_key, year if year >= 0 else None, quarter_label


def table_domain(pb_table: int) -> str:
    if pb_table in {1, 2, 3}:
        return "bop-macro"
    if pb_table == 23:
        return "services"
    if pb_table in {25, 26}:
        return "income-transfers"
    if pb_table in {28, 29, 31, 32}:
        return "investment-financial"
    if pb_table in {4, 5, 6, 24, 27, 30}:
        return "trade-geography"
    if 7 <= pb_table <= 22:
        return "commodity-trade"
    return "other"


def normalize_header(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip().lower().replace(" ", "_")


def column_index_to_letter(index: int) -> str:
    letters = ""
    while index > 0:
        index, remainder = divmod(index - 1, 26)
        letters = chr(65 + remainder) + letters
    return letters


def column_letter_to_index(letter: str) -> int:
    result = 0
    for ch in letter.upper():
        if not ("A" <= ch <= "Z"):
            raise ValueError(f"Invalid column letter: {letter}")
        result = result * 26 + (ord(ch) - 64)
    return result


def hierarchy_level(code: str | None) -> int | None:
    if not code:
        return None
    value = code.strip()
    if not value:
        return None
    if re.match(r"^[IVX]+\.$", value):
        return 1
    if re.match(r"^\d+(?:\.\d+)*\.?$", value):
        trimmed = value.rstrip(".")
        return trimmed.count(".") + 1
    if "." in value:
        return value.count(".") + 1
    return 1


def parent_code(code: str | None) -> str | None:
    if not code:
        return None
    value = code.strip()
    if not value or "." not in value:
        return None
    trimmed = value.rstrip(".")
    if "." not in trimmed:
        return None
    parent = trimmed.rsplit(".", 1)[0]
    return parent + "."


def choose_name_columns(header_map: dict[str, int]) -> dict[str, int]:
    resolved: dict[str, int] = {}
    preference = {
        "name_tg": ["names_tg", "item_name_taj"],
        "name_ru": ["names_ru", "item_name_ru"],
        "name_en": ["names_en", "item_name_en"],
        "name_primary": ["names", "item_name", "item_name_en", "names_en"],
    }
    for target, keys in preference.items():
        for key in keys:
            if key in header_map:
                resolved[target] = header_map[key]
                break
    return resolved


def infer_measure_columns(
    header_map: dict[str, int],
    ws: Any,
    config_value_cols: dict[str, str | dict[str, str]] | None,
) -> tuple[dict[str, int], str, str]:
    excluded = {
        "item_code",
        "code",
        "code_guid",
        "names",
        "names_tg",
        "names_ru",
        "names_en",
        "item_name_taj",
        "item_name_ru",
        "item_name_en",
        "year",
        "quarter",
        "period",
        "tnved_code",
        "product_code",
        "gid_code",
        "data_type",
        "item_name",
    }

    def numeric_count_for_column(col_index: int) -> int:
        count = 0
        for row in range(2, min(ws.max_row, 250) + 1):
            value = ws.cell(row=row, column=col_index).value
            if isinstance(value, (int, float)):
                count += 1
        return count

    if config_value_cols:
        resolved: dict[str, int] = {}
        for measure_name, descriptor in config_value_cols.items():
            if isinstance(descriptor, str):
                try:
                    resolved[measure_name] = column_letter_to_index(descriptor)
                except ValueError:
                    continue
            elif isinstance(descriptor, dict):
                column_letter = descriptor.get("column")
                if column_letter:
                    try:
                        resolved[measure_name] = column_letter_to_index(column_letter)
                    except ValueError:
                        continue
        # Validate config columns against actual workbook shape. Some configs map
        # legacy templates where column letters no longer match extracted files.
        validated: dict[str, int] = {}
        rejection_reasons: list[str] = []
        index_to_header = {
            idx: header for header, idx in header_map.items()
        }
        for measure_name, col_index in resolved.items():
            if col_index > ws.max_column:
                rejection_reasons.append(
                    f"{measure_name}: column index {col_index} is out of range"
                )
                continue
            header_name = index_to_header.get(col_index, "")
            if header_name in excluded:
                rejection_reasons.append(
                    f"{measure_name}: mapped to excluded header '{header_name}'"
                )
                continue
            numeric_count = numeric_count_for_column(col_index)
            if numeric_count < 3:
                rejection_reasons.append(
                    f"{measure_name}: column {col_index} has only {numeric_count} numeric rows"
                )
                continue
            validated[measure_name] = col_index
        if validated:
            return validated, "config", ""
        fallback_reason = (
            "; ".join(rejection_reasons)
            if rejection_reasons
            else "configured value_cols are missing or invalid for workbook shape"
        )
    else:
        fallback_reason = ""

    numeric_candidates: dict[str, int] = {}
    for header, col_index in header_map.items():
        if header in excluded or not header:
            continue
        numeric_count = numeric_count_for_column(col_index)
        if numeric_count >= 3:
            numeric_candidates[header] = col_index
    if config_value_cols:
        return numeric_candidates, "fallback", fallback_reason
    return numeric_candidates, "inferred", ""


@dataclass
class CategorySeen:
    domain: str
    pb_table: int
    category_code: str
    category_name_primary: str
    category_name_tg: str
    category_name_ru: str
    category_name_en: str
    hierarchy_level: int | None
    first_seen_period: str
    first_seen_sort_key: int
    last_seen_period: str
    last_seen_sort_key: int
    last_seen_file: str
    seen_files: set[str]


def write_csv(path: Path, rows: list[dict[str, Any]], fieldnames: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def main() -> int:
    args = parse_args()

    output_rar = Path(args.output_rar).resolve()
    config_rar = Path(args.config_rar).resolve()
    output_dir = Path(args.output_dir).resolve()
    work_dir = Path(args.work_dir).resolve()
    extract_output_dir = work_dir / "output_files_extracted"
    extract_config_dir = work_dir / "report_configs_extracted"
    logs_dir = output_dir / "_logs"

    if not args.skip_extract:
        require_utility("unar")
        if work_dir.exists() and not args.keep_work_dir:
            shutil.rmtree(work_dir)
        work_dir.mkdir(parents=True, exist_ok=True)
        extract_output_dir.mkdir(parents=True, exist_ok=True)
        extract_config_dir.mkdir(parents=True, exist_ok=True)

        if not output_rar.exists():
            raise FileNotFoundError(f"Archive not found: {output_rar}")
        if not config_rar.exists():
            raise FileNotFoundError(f"Archive not found: {config_rar}")

        output_rc, output_log = run_unar(output_rar, extract_output_dir)
        config_rc, config_log = run_unar(config_rar, extract_config_dir)

        logs_dir.mkdir(parents=True, exist_ok=True)
        (logs_dir / "extract_output_files.log").write_text(
            output_log, encoding="utf-8"
        )
        (logs_dir / "extract_report_configs.log").write_text(
            config_log, encoding="utf-8"
        )

        # unar may return non-zero when some files fail; continue with available files.
        if output_rc != 0:
            print(
                "WARN: output_files extraction returned non-zero; "
                "continuing with extracted files."
            )
        if config_rc != 0:
            print(
                "WARN: report_configs extraction returned non-zero; "
                "continuing with extracted files."
            )
    else:
        if not extract_output_dir.exists() or not extract_config_dir.exists():
            raise RuntimeError(
                "--skip-extract was set, but extracted directories are missing:\n"
                f"- {extract_output_dir}\n- {extract_config_dir}"
            )

    cfg_root = extract_config_dir / "report_configs"
    out_root = extract_output_dir / "output_files"

    if not cfg_root.exists():
        raise RuntimeError(f"Missing extracted config folder: {cfg_root}")
    if not out_root.exists():
        raise RuntimeError(f"Missing extracted workbook folder: {out_root}")

    config_by_pb: dict[int, dict[str, Any]] = {}
    config_files = sorted(cfg_root.glob("pb_t*.json"))
    for cfg_file in config_files:
        match = re.match(r"pb_t(\d+)", cfg_file.name)
        if not match:
            continue
        pb_table = int(match.group(1))
        if cfg_file.stat().st_size == 0:
            continue
        try:
            config_by_pb[pb_table] = json.loads(cfg_file.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            continue

    workbook_files = sorted(out_root.glob("PB*.xlsx"))

    dim_report_period: dict[str, dict[str, Any]] = {}
    dim_measure_type: dict[str, dict[str, Any]] = {}
    dim_data_scope: dict[str, dict[str, Any]] = {}
    dim_table: dict[int, dict[str, Any]] = {}
    category_seen: dict[
        tuple[int, str, str, str, str, str], CategorySeen
    ] = {}
    category_key_lookup: dict[tuple[int, str], str] = {}

    fact_rows: list[dict[str, Any]] = []
    fallback_rows: list[dict[str, Any]] = []
    data_type_last_seen: dict[tuple[int, str, str], dict[str, Any]] = {}

    for pb_table in range(1, 33):
        dim_table[pb_table] = {
            "table_key": f"PB{pb_table}",
            "pb_table": pb_table,
            "table_code": f"PB{pb_table}",
            "table_domain": table_domain(pb_table),
        }

    for file_path in workbook_files:
        match = re.match(r"PB(\d+)\.", file_path.name)
        if not match:
            continue
        pb_table = int(match.group(1))
        if file_path.stat().st_size == 0:
            continue

        report_label, report_sort_key, report_year, report_quarter = (
            parse_period_from_filename(file_path.name)
        )
        dim_report_period[report_label] = {
            "report_period_key": report_label,
            "report_period_label": report_label,
            "report_year": report_year if report_year is not None else "",
            "report_quarter": report_quarter,
            "period_sort_key": report_sort_key,
        }

        try:
            wb = load_workbook(file_path, data_only=True)
            ws = wb[wb.sheetnames[0]]
        except Exception:
            continue

        headers = [ws.cell(row=1, column=i).value for i in range(1, ws.max_column + 1)]
        header_map: dict[str, int] = {}
        for index, raw_header in enumerate(headers, start=1):
            normalized = normalize_header(raw_header)
            if normalized:
                header_map[normalized] = index

        code_col = None
        for candidate in ("item_code", "code"):
            if candidate in header_map:
                code_col = header_map[candidate]
                break
        if code_col is None:
            for key, value in header_map.items():
                if "code" in key and "tnved" not in key and "gid" not in key:
                    code_col = value
                    break

        name_cols = choose_name_columns(header_map)
        primary_name_col = name_cols.get("name_primary") or name_cols.get("name_en")
        if primary_name_col is None:
            # fallback to any name-like column
            for key, value in header_map.items():
                if "name" in key or "names" in key:
                    primary_name_col = value
                    break

        period_col = header_map.get("period")
        data_type_col = header_map.get("data_type")
        tnved_col = header_map.get("tnved_code")
        product_col = header_map.get("product_code")
        year_col = header_map.get("year")
        quarter_col = header_map.get("quarter")

        cfg = config_by_pb.get(pb_table, {})
        cfg_value_cols = (cfg.get("columns", {}) or {}).get("value_cols")
        measure_cols, measure_strategy, fallback_reason = infer_measure_columns(
            header_map, ws, cfg_value_cols
        )
        if measure_strategy == "fallback":
            configured_names = (
                ",".join(sorted(cfg_value_cols.keys()))
                if isinstance(cfg_value_cols, dict)
                else ""
            )
            inferred_names = ",".join(sorted(measure_cols.keys()))
            message = (
                "WARN: fallback to inferred measure columns for "
                f"{file_path.name} (PB{pb_table}). "
                f"configured=[{configured_names}] inferred=[{inferred_names}]"
            )
            print(message)
            fallback_rows.append(
                {
                    "pb_table": pb_table,
                    "source_file": file_path.name,
                    "source_sheet": wb.sheetnames[0],
                    "configured_measures": configured_names,
                    "inferred_measures": inferred_names,
                    "reason": fallback_reason,
                }
            )
        if not measure_cols:
            continue

        # register config data scopes
        for data_type in cfg.get("data_types", []) or []:
            scope_code = data_type.get("name")
            if not scope_code:
                continue
            dim_data_scope[scope_code] = {
                "data_scope_key": scope_code,
                "scope_code": scope_code,
                "label_tg": data_type.get("label_tg", ""),
                "label_ru": data_type.get("label_ru", ""),
                "label_en": data_type.get("label_en", ""),
            }
            data_type_last_seen[(pb_table, "data_type", scope_code)] = {
                "pb_table": pb_table,
                "dimension_type": "data_type",
                "category_key": scope_code,
                "label_tg": data_type.get("label_tg", ""),
                "label_ru": data_type.get("label_ru", ""),
                "label_en": data_type.get("label_en", ""),
                "last_seen_period": report_label,
                "last_seen_sort_key": report_sort_key,
                "last_seen_file": file_path.name,
            }

        for row_index in range(2, ws.max_row + 1):
            code_value = (
                ws.cell(row=row_index, column=code_col).value if code_col else None
            )
            primary_name_value = (
                ws.cell(row=row_index, column=primary_name_col).value
                if primary_name_col
                else None
            )
            if code_value in (None, "") and primary_name_value in (None, ""):
                continue

            category_code = "" if code_value is None else str(code_value).strip()
            category_name_primary = (
                "" if primary_name_value is None else str(primary_name_value).strip()
            )
            category_name_tg = (
                ""
                if "name_tg" not in name_cols
                else str(ws.cell(row=row_index, column=name_cols["name_tg"]).value or "").strip()
            )
            category_name_ru = (
                ""
                if "name_ru" not in name_cols
                else str(ws.cell(row=row_index, column=name_cols["name_ru"]).value or "").strip()
            )
            category_name_en = (
                ""
                if "name_en" not in name_cols
                else str(ws.cell(row=row_index, column=name_cols["name_en"]).value or "").strip()
            )

            if not category_name_en and category_name_primary:
                category_name_en = category_name_primary

            level = hierarchy_level(category_code)
            cat_key = (
                pb_table,
                category_code,
                category_name_primary,
                category_name_tg,
                category_name_ru,
                category_name_en,
            )
            seen = category_seen.get(cat_key)
            if seen is None:
                seen = CategorySeen(
                    domain=table_domain(pb_table),
                    pb_table=pb_table,
                    category_code=category_code,
                    category_name_primary=category_name_primary,
                    category_name_tg=category_name_tg,
                    category_name_ru=category_name_ru,
                    category_name_en=category_name_en,
                    hierarchy_level=level,
                    first_seen_period=report_label,
                    first_seen_sort_key=report_sort_key,
                    last_seen_period=report_label,
                    last_seen_sort_key=report_sort_key,
                    last_seen_file=file_path.name,
                    seen_files={file_path.name},
                )
                category_seen[cat_key] = seen
            else:
                if report_sort_key < seen.first_seen_sort_key:
                    seen.first_seen_sort_key = report_sort_key
                    seen.first_seen_period = report_label
                if report_sort_key >= seen.last_seen_sort_key:
                    seen.last_seen_sort_key = report_sort_key
                    seen.last_seen_period = report_label
                    seen.last_seen_file = file_path.name
                seen.seen_files.add(file_path.name)

            category_member_key = (
                f"PB{pb_table}|{category_code}|{category_name_primary or category_name_en}"
            )
            if category_code:
                category_key_lookup[(pb_table, category_code)] = category_member_key

            period_bucket = ""
            if period_col:
                period_value = ws.cell(row=row_index, column=period_col).value
                if period_value is not None:
                    period_bucket = str(period_value).strip()

            data_scope = ""
            if data_type_col:
                scope_value = ws.cell(row=row_index, column=data_type_col).value
                if scope_value is not None:
                    data_scope = str(scope_value).strip()
                    if data_scope and data_scope not in dim_data_scope:
                        dim_data_scope[data_scope] = {
                            "data_scope_key": data_scope,
                            "scope_code": data_scope,
                            "label_tg": "",
                            "label_ru": "",
                            "label_en": "",
                        }

            source_year = (
                str(ws.cell(row=row_index, column=year_col).value).strip()
                if year_col and ws.cell(row=row_index, column=year_col).value is not None
                else ""
            )
            source_quarter = (
                str(ws.cell(row=row_index, column=quarter_col).value).strip()
                if quarter_col
                and ws.cell(row=row_index, column=quarter_col).value is not None
                else ""
            )

            tnved_code = (
                str(ws.cell(row=row_index, column=tnved_col).value).strip()
                if tnved_col and ws.cell(row=row_index, column=tnved_col).value is not None
                else ""
            )
            product_code = (
                str(ws.cell(row=row_index, column=product_col).value).strip()
                if product_col
                and ws.cell(row=row_index, column=product_col).value is not None
                else ""
            )

            for measure_name, measure_col in measure_cols.items():
                if measure_col > ws.max_column:
                    continue
                value = ws.cell(row=row_index, column=measure_col).value
                if value in (None, ""):
                    continue
                if not isinstance(value, (int, float)):
                    continue
                effective_measure_name = measure_name
                if (
                    measure_name in {"value", "values"}
                    and period_bucket
                ):
                    effective_measure_name = period_bucket

                fact_rows.append(
                    {
                        "table_key": f"PB{pb_table}",
                        "pb_table": pb_table,
                        "report_period_key": report_label,
                        "source_year_value": source_year,
                        "source_quarter_value": source_quarter,
                        "category_member_key": category_member_key,
                        "category_code": category_code,
                        "measure_type_key": effective_measure_name,
                        "data_scope_key": data_scope,
                        "period_bucket": period_bucket,
                        "tnved_code": tnved_code,
                        "product_code": product_code,
                        "source_file": file_path.name,
                        "source_sheet": wb.sheetnames[0],
                        "source_row": row_index,
                        "source_column_letter": column_index_to_letter(measure_col),
                        "measure_value": value,
                    }
                )

    period_enforcement_rows = enforce_current_year_for_sparse_periods(fact_rows)

    # Rebuild measure dim after potential period normalization.
    dim_measure_type = {}
    for measure_type_key in sorted(
        {str(row.get("measure_type_key") or "").strip() for row in fact_rows if str(row.get("measure_type_key") or "").strip()}
    ):
        dim_measure_type[measure_type_key] = {
            "measure_type_key": measure_type_key,
            "measure_code": measure_type_key,
            "measure_family": measure_type_key,
        }

    # Build dim_category_member
    dim_category_rows: list[dict[str, Any]] = []
    for key, seen in category_seen.items():
        category_member_key = (
            f"PB{seen.pb_table}|{seen.category_code}|"
            f"{seen.category_name_primary or seen.category_name_en}"
        )
        dim_category_rows.append(
            {
                "category_member_key": category_member_key,
                "domain": seen.domain,
                "pb_table": seen.pb_table,
                "category_code": seen.category_code,
                "category_name_primary": seen.category_name_primary,
                "category_name_tg": seen.category_name_tg,
                "category_name_ru": seen.category_name_ru,
                "category_name_en": seen.category_name_en,
                "hierarchy_level": seen.hierarchy_level if seen.hierarchy_level else "",
                "parent_category_code": parent_code(seen.category_code) or "",
                "first_seen_period_key": seen.first_seen_period,
                "last_seen_period_key": seen.last_seen_period,
                "last_seen_file": seen.last_seen_file,
                "seen_in_files": len(seen.seen_files),
            }
        )

    dim_category_rows.sort(
        key=lambda row: (
            row["pb_table"],
            row["category_code"],
            row["category_name_primary"],
        )
    )

    # Build closure bridge
    # pick deterministic category key by pb/code using first matching row
    pb_code_to_key: dict[tuple[int, str], str] = {}
    for row in dim_category_rows:
        code = row["category_code"]
        if code and (row["pb_table"], code) not in pb_code_to_key:
            pb_code_to_key[(row["pb_table"], code)] = row["category_member_key"]

    closure_rows: list[dict[str, Any]] = []
    for row in dim_category_rows:
        descendant_key = row["category_member_key"]
        pb_table = row["pb_table"]
        code = row["category_code"]
        if not code:
            continue
        # self edge
        closure_rows.append(
            {
                "ancestor_category_member_key": descendant_key,
                "descendant_category_member_key": descendant_key,
                "distance": 0,
            }
        )
        current_code = code
        distance = 1
        while True:
            parent = parent_code(current_code)
            if not parent:
                break
            ancestor_key = pb_code_to_key.get((pb_table, parent))
            if not ancestor_key:
                break
            closure_rows.append(
                {
                    "ancestor_category_member_key": ancestor_key,
                    "descendant_category_member_key": descendant_key,
                    "distance": distance,
                }
            )
            current_code = parent
            distance += 1

    # Build reference table rows
    ref_category_rows: list[dict[str, Any]] = []
    for key, seen in category_seen.items():
        ref_category_rows.append(
            {
                "domain": seen.domain,
                "pb_table": seen.pb_table,
                "category_code": seen.category_code,
                "category_name_primary": seen.category_name_primary,
                "category_name_tg": seen.category_name_tg,
                "category_name_ru": seen.category_name_ru,
                "category_name_en": seen.category_name_en,
                "hierarchy_level": seen.hierarchy_level if seen.hierarchy_level else "",
                "first_seen_period": seen.first_seen_period,
                "last_seen_period": seen.last_seen_period,
                "last_seen_file": seen.last_seen_file,
                "seen_in_files": len(seen.seen_files),
            }
        )
    ref_category_rows.sort(
        key=lambda row: (
            row["domain"],
            row["pb_table"],
            row["category_code"],
            row["category_name_primary"],
        )
    )

    period_sort_lookup = {
        key: int(value["period_sort_key"])
        for key, value in dim_report_period.items()
    }

    period_bucket_seen: dict[tuple[int, str], dict[str, Any]] = {}
    for row in fact_rows:
        period_bucket = str(row.get("period_bucket") or "").strip()
        if period_bucket not in PERIOD_BUCKETS:
            continue
        pb_table = int(row.get("pb_table") or 0)
        report_period_key = str(row.get("report_period_key") or "")
        report_sort_key = period_sort_lookup.get(report_period_key, -1)
        source_file = str(row.get("source_file") or "")
        key = (pb_table, period_bucket)
        existing = period_bucket_seen.get(key)
        if existing is None:
            period_bucket_seen[key] = {
                "pb_table": pb_table,
                "period_bucket": period_bucket,
                "first_seen_period": report_period_key,
                "first_seen_sort_key": report_sort_key,
                "last_seen_period": report_period_key,
                "last_seen_sort_key": report_sort_key,
                "seen_in_files": {source_file} if source_file else set(),
            }
        else:
            if report_sort_key < existing["first_seen_sort_key"]:
                existing["first_seen_sort_key"] = report_sort_key
                existing["first_seen_period"] = report_period_key
            if report_sort_key >= existing["last_seen_sort_key"]:
                existing["last_seen_sort_key"] = report_sort_key
                existing["last_seen_period"] = report_period_key
            if source_file:
                existing["seen_in_files"].add(source_file)

    ref_period_rows: list[dict[str, Any]] = []
    for _, item in sorted(period_bucket_seen.items()):
        ref_period_rows.append(
            {
                "pb_table": item["pb_table"],
                "period_bucket": item["period_bucket"],
                "first_seen_period": item["first_seen_period"],
                "last_seen_period": item["last_seen_period"],
                "seen_in_files": len(item["seen_in_files"]),
            }
        )

    value_type_last_seen: dict[tuple[int, str, str], dict[str, Any]] = {}
    for row in fact_rows:
        pb_table = int(row.get("pb_table") or 0)
        measure_type_key = str(row.get("measure_type_key") or "").strip()
        if not measure_type_key:
            continue
        report_period_key = str(row.get("report_period_key") or "")
        report_sort_key = period_sort_lookup.get(report_period_key, -1)
        source_file = str(row.get("source_file") or "")
        key = (pb_table, "value_column", measure_type_key)
        current = value_type_last_seen.get(key)
        if current is None or report_sort_key >= current["last_seen_sort_key"]:
            value_type_last_seen[key] = {
                "pb_table": pb_table,
                "dimension_type": "value_column",
                "category_key": measure_type_key,
                "label_tg": "",
                "label_ru": "",
                "label_en": "",
                "last_seen_period": report_period_key,
                "last_seen_sort_key": report_sort_key,
                "last_seen_file": source_file,
            }

    ref_type_rows: list[dict[str, Any]] = []
    combined_type_last_seen = dict(data_type_last_seen)
    combined_type_last_seen.update(value_type_last_seen)
    for _, item in sorted(combined_type_last_seen.items()):
        ref_type_rows.append(
            {
                "pb_table": item["pb_table"],
                "dimension_type": item["dimension_type"],
                "category_key": item["category_key"],
                "label_tg": item["label_tg"],
                "label_ru": item["label_ru"],
                "label_en": item["label_en"],
                "last_seen_period": item["last_seen_period"],
                "last_seen_file": item["last_seen_file"],
            }
        )

    output_dir.mkdir(parents=True, exist_ok=True)
    logs_dir.mkdir(parents=True, exist_ok=True)

    write_csv(
        output_dir / "dim_table.csv",
        [v for _, v in sorted(dim_table.items())],
        ["table_key", "pb_table", "table_code", "table_domain"],
    )
    write_csv(
        output_dir / "dim_report_period.csv",
        sorted(dim_report_period.values(), key=lambda row: row["period_sort_key"]),
        [
            "report_period_key",
            "report_period_label",
            "report_year",
            "report_quarter",
            "period_sort_key",
        ],
    )
    write_csv(
        output_dir / "dim_measure_type.csv",
        sorted(dim_measure_type.values(), key=lambda row: row["measure_type_key"]),
        ["measure_type_key", "measure_code", "measure_family"],
    )
    write_csv(
        output_dir / "dim_data_scope.csv",
        sorted(dim_data_scope.values(), key=lambda row: row["scope_code"]),
        ["data_scope_key", "scope_code", "label_tg", "label_ru", "label_en"],
    )
    write_csv(
        output_dir / "dim_category_member.csv",
        dim_category_rows,
        [
            "category_member_key",
            "domain",
            "pb_table",
            "category_code",
            "category_name_primary",
            "category_name_tg",
            "category_name_ru",
            "category_name_en",
            "hierarchy_level",
            "parent_category_code",
            "first_seen_period_key",
            "last_seen_period_key",
            "last_seen_file",
            "seen_in_files",
        ],
    )
    write_csv(
        output_dir / "bridge_category_closure.csv",
        closure_rows,
        [
            "ancestor_category_member_key",
            "descendant_category_member_key",
            "distance",
        ],
    )
    write_csv(
        output_dir / "fact_rows_long.csv",
        fact_rows,
        [
            "table_key",
            "pb_table",
            "report_period_key",
            "source_year_value",
            "source_quarter_value",
            "category_member_key",
            "category_code",
            "measure_type_key",
            "data_scope_key",
            "period_bucket",
            "tnved_code",
            "product_code",
            "source_file",
            "source_sheet",
            "source_row",
            "source_column_letter",
            "measure_value",
        ],
    )

    write_csv(
        output_dir / "ref_dim_category_last_seen.csv",
        ref_category_rows,
        [
            "domain",
            "pb_table",
            "category_code",
            "category_name_primary",
            "category_name_tg",
            "category_name_ru",
            "category_name_en",
            "hierarchy_level",
            "first_seen_period",
            "last_seen_period",
            "last_seen_file",
            "seen_in_files",
        ],
    )
    write_csv(
        output_dir / "ref_dim_period_bucket_last_seen.csv",
        ref_period_rows,
        [
            "pb_table",
            "period_bucket",
            "first_seen_period",
            "last_seen_period",
            "seen_in_files",
        ],
    )
    write_csv(
        output_dir / "ref_dim_type_last_seen.csv",
        ref_type_rows,
        [
            "pb_table",
            "dimension_type",
            "category_key",
            "label_tg",
            "label_ru",
            "label_en",
            "last_seen_period",
            "last_seen_file",
        ],
    )
    write_csv(
        logs_dir / "measure_column_fallbacks.csv",
        fallback_rows,
        [
            "pb_table",
            "source_file",
            "source_sheet",
            "configured_measures",
            "inferred_measures",
            "reason",
        ],
    )
    write_csv(
        logs_dir / "period_bucket_enforcement.csv",
        period_enforcement_rows,
        [
            "pb_table",
            "original_period_buckets",
            "enforced_period_bucket",
        ],
    )

    summary = {
        "files_processed": len([f for f in workbook_files if f.stat().st_size > 0]),
        "facts_rows": len(fact_rows),
        "dim_category_rows": len(dim_category_rows),
        "period_rows": len(dim_report_period),
        "measure_types": len(dim_measure_type),
        "data_scopes": len(dim_data_scope),
        "measure_column_fallbacks": len(fallback_rows),
        "period_bucket_enforcements": len(period_enforcement_rows),
    }
    (output_dir / "run_summary.json").write_text(
        json.dumps(summary, indent=2), encoding="utf-8"
    )

    print("Prepared model pieces written to:", output_dir)
    print(json.dumps(summary, indent=2))
    return 0


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except Exception as exc:  # pragma: no cover
        print(f"ERROR: {exc}", file=sys.stderr)
        raise
